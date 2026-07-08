from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path
from typing import Any, Iterable


REPORT_PRIORITY_COLUMNS = [
    "report_rank",
    "rank_bucket",
    "rank_bucket_label",
    "report_priority_score",
    "identity_preference",
    "recommendation",
    "candidate_id",
    "gene_symbol",
    "uniprot_accession",
    "protein_name",
    "membrane_group",
    "antigen_accessibility",
    "accessibility_gate",
    "ecd_region",
    "ecd_length",
    "human_mouse_ecd_identity",
    "human_mouse_ecd_identity_pct",
    "human_mouse_divergence_class",
    "original_mouse_model_transferability_gate",
    "constructability_gate",
    "constructability_class",
    "construct_recommendation",
    "primary_strategy",
    "secondary_strategy",
    "normal_tissue_risk",
    "normal_tissue_risk_detail",
    "pipeline_priority_call",
    "local_evidence_source_count",
    "local_evidence_sources",
    "risk_flags_original",
    "readiness_flags",
    "passed_reasons",
    "caution_or_fail_reasons",
    "next_action",
    "coverage_missing_reason",
]


SOURCE_COLUMNS = [
    ("evidence_2021_bioinfo_wang", "2021_bioinfo_wang"),
    ("evidence_2008_nmeth_labeer", "2008_nmeth_labeer"),
    ("evidence_igdomain_prot_list", "Ig_domain_list"),
    ("evidence_reap_list", "REAP"),
    ("evidence_fda_carttarget_list", "FDA_CAR_T"),
    ("evidence_2024_cell_drugtarget", "2024_Cell_drugtarget"),
    ("evidence_drugtarget_prot_list", "DrugTarget_protein"),
    ("evidence_ohs6087_ccsb_list", "OHS6087_CCSB"),
    ("evidence_local_prot_list", "Local_protein"),
]


def _safe_float(value: object) -> float | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        parsed = float(text)
    except ValueError:
        return None
    return parsed / 100.0 if parsed > 1 else parsed


def _safe_int(value: object) -> int | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _pct(value: float | None) -> str:
    return "" if value is None else f"{value * 100:.1f}%"


def _read_optional(path: Path | None) -> list[dict[str, str]]:
    if not path or not path.exists() or path.stat().st_size == 0:
        return []
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def _by_id(rows: Iterable[dict[str, str]]) -> dict[str, dict[str, str]]:
    return {row.get("candidate_id", ""): row for row in rows if row.get("candidate_id")}


def _construct_by_gene(rows: Iterable[dict[str, str]]) -> dict[str, dict[str, str]]:
    return {row.get("gene_symbol", ""): row for row in rows if row.get("gene_symbol")}


def _split_flags(*values: object) -> list[str]:
    seen: list[str] = []
    for value in values:
        for part in str(value or "").replace("|", ";").split(";"):
            item = part.strip()
            if item and item not in seen:
                seen.append(item)
    return seen


def _divergence_class(identity: float | None, preference: str) -> tuple[str, float]:
    if identity is None:
        return "missing_identity", 0
    if preference == "high_for_mouse_model_transferability":
        if identity >= 0.90:
            return "high_conservation_priority", 80
        if identity >= 0.80:
            return "moderate_conservation", 55
        if identity >= 0.70:
            return "limited_conservation", 25
        return "low_conservation_model_risk", 0
    if identity < 0.70:
        return "high_divergence_priority", 80
    if identity < 0.80:
        return "moderate_divergence", 55
    if identity < 0.90:
        return "limited_divergence", 25
    return "highly_conserved_lower_antibody_priority", 0


def rank_candidates(
    screening_rows: list[dict[str, str]],
    *,
    readiness_rows: list[dict[str, str]] | None = None,
    normal_tissue_rows: list[dict[str, str]] | None = None,
    construct_rows: list[dict[str, str]] | None = None,
    strategy_rows: list[dict[str, str]] | None = None,
    coverage_rows: list[dict[str, str]] | None = None,
    join_rows: list[dict[str, str]] | None = None,
    identity_preference: str = "low_for_antibody_screening",
) -> list[dict[str, Any]]:
    if identity_preference not in {"low_for_antibody_screening", "high_for_mouse_model_transferability"}:
        raise ValueError(f"unsupported identity_preference: {identity_preference}")

    readiness_by_id = _by_id(readiness_rows or [])
    normal_by_id = _by_id(normal_tissue_rows or [])
    strategy_by_id = _by_id(strategy_rows or [])
    coverage_by_id = _by_id(coverage_rows or [])
    join_by_id = _by_id(join_rows or [])
    construct_by_gene = _construct_by_gene(construct_rows or [])

    ranked: list[dict[str, Any]] = []
    for row in screening_rows:
        candidate_id = row.get("candidate_id", "")
        gene_symbol = row.get("gene_symbol", "")
        readiness = readiness_by_id.get(candidate_id, {})
        normal = normal_by_id.get(candidate_id, {})
        strategy = strategy_by_id.get(candidate_id, {})
        coverage = coverage_by_id.get(candidate_id, {})
        join = join_by_id.get(candidate_id, {})
        construct = construct_by_gene.get(gene_symbol, {})

        identity = _safe_float(row.get("human_mouse_ecd_identity"))
        divergence_class, divergence_score = _divergence_class(identity, identity_preference)
        ecd_length = _safe_int(row.get("ecd_length"))
        length_ok = ecd_length is not None and 40 <= ecd_length <= 1000
        length_optimal = ecd_length is not None and 80 <= ecd_length <= 800
        accessibility = row.get("antigen_accessibility", "")
        access_gate = row.get("accessibility_gate", "")
        construct_gate = row.get("constructability_gate", "")
        normal_risk = normal.get("normal_tissue_risk") or row.get("normal_tissue_risk") or "unknown"
        high_risk_tissues = normal.get("high_risk_tissue_flags", "")
        surface_clear = accessibility in {"cell_surface_single_pass", "gpi_anchor", "surface_single_pass", "gpi_anchored"}
        mixed_or_secreted = accessibility in {"surface_or_secreted_mixed_annotation", "secreted_soluble", "secreted"}
        local_sources = [
            label
            for column, label in SOURCE_COLUMNS
            if str(join.get(column, "")).strip()
        ]
        risk_flags = _split_flags(row.get("risk_flags"))
        readiness_flags = _split_flags(readiness.get("readiness_flags"))

        if access_gate == "fail" or normal_risk == "high":
            bucket = "D"
            bucket_label = "D_deprioritize_or_verify"
            recommendation = "deprioritize_for_main_screen_until_risk_or_accessibility_is_resolved"
        elif (
            identity_preference == "low_for_antibody_screening"
            and construct_gate == "pass"
            and surface_clear
            and identity is not None
            and identity < 0.80
            and length_ok
        ):
            bucket = "A"
            bucket_label = "A_priority_expression_screen"
            recommendation = "prioritize_for_ecd_expression_and_antibody_screen"
        elif construct_gate == "pass" and identity is not None:
            bucket = "B"
            bucket_label = "B_screenable_needs_risk_or_identity_review"
            recommendation = "screenable_but_complete_normal_tissue_and_identity_review"
        elif construct_gate == "partial" or row.get("priority_call") == "conditional":
            bucket = "C"
            bucket_label = "C_conditional_construct_or_format_optimization"
            recommendation = "optimize_construct_or_display_before_main_screen"
        else:
            bucket = "C"
            bucket_label = "C_conditional_evidence_incomplete"
            recommendation = "hold_for_evidence_completion_or_secondary_screen"

        score = {"A": 400, "B": 300, "C": 180, "D": 0}[bucket]
        score += divergence_score
        score += 35 if construct_gate == "pass" else 8 if construct_gate == "partial" else -20
        if surface_clear:
            score += 30
        elif accessibility == "surface_or_secreted_mixed_annotation":
            score += 12
        elif accessibility in {"secreted_soluble", "secreted"}:
            score -= 10
        elif accessibility in {"cell_surface_multi_pass", "surface_multi_pass"}:
            score -= 15
        score += 18 if length_optimal else 8 if length_ok else -12
        score += min(len(local_sources), 8) * 3
        if normal_risk == "unknown":
            score -= 8
        elif normal_risk == "high":
            score -= 90
        if "secreted_only" in risk_flags:
            score -= 12
        if "very_long_ecd" in readiness_flags:
            score -= 12
        if "short_ecd" in readiness_flags:
            score -= 8
        if "isoform_ambiguity" in readiness_flags or row.get("isoform_risk") == "missing_live_isoform_review":
            score -= 5
        if "fda_cart_evidence_context" in risk_flags:
            score -= 10

        passed_reasons = []
        if access_gate == "pass":
            passed_reasons.append("antigen_accessibility_pass")
        if surface_clear:
            passed_reasons.append("clear_surface_or_gpi_annotation")
        elif accessibility == "surface_or_secreted_mixed_annotation":
            passed_reasons.append("mixed_surface_secreted_annotation")
        elif accessibility in {"secreted_soluble", "secreted"}:
            passed_reasons.append("soluble_secreted_antigen_screenable")
        if identity is not None:
            passed_reasons.append(f"human_mouse_ecd_identity={_pct(identity)}:{divergence_class}")
        if construct_gate == "pass":
            passed_reasons.append("cleaner_soluble_ecd_construct")
        elif construct_gate == "partial":
            passed_reasons.append("construct_possible_with_special_strategy")
        if local_sources:
            passed_reasons.append(f"local_evidence_sources={len(local_sources)}")

        cautions = []
        if normal_risk == "unknown":
            cautions.append("normal_tissue_live_hpa_gtex_missing")
        elif normal_risk == "high":
            cautions.append(f"high_normal_tissue_risk:{high_risk_tissues or 'high'}")
        if identity is None:
            cautions.append("human_mouse_ecd_identity_missing")
        elif identity_preference == "low_for_antibody_screening" and identity >= 0.90:
            cautions.append("highly_conserved_not_low_identity_priority")
        elif identity_preference == "low_for_antibody_screening" and identity >= 0.80:
            cautions.append("limited_human_mouse_ecd_divergence")
        if accessibility == "surface_or_secreted_mixed_annotation":
            cautions.append("verify_surface_exposure_and_isoform")
        if accessibility in {"secreted_soluble", "secreted"}:
            cautions.append("secreted_not_typical_cell_surface_target")
        if "very_long_ecd" in readiness_flags:
            cautions.append("very_long_ecd_expression_risk")
        if "short_ecd" in readiness_flags:
            cautions.append("short_ecd_may_need_multimer_or_fragment_strategy")
        if row.get("isoform_risk") == "missing_live_isoform_review" or "isoform_ambiguity" in readiness_flags:
            cautions.append("live_isoform_topology_review_pending")
        if coverage.get("missing_reason_primary"):
            cautions.append(f"coverage_gap:{coverage.get('missing_reason_primary')}")

        next_actions = []
        if bucket == "A":
            next_actions.append("start_small_scale_ecd_or_fc_his_avi_expression")
        elif bucket == "B":
            next_actions.append("screenable_after_surface_and_normal_tissue_review")
        elif bucket == "C":
            next_actions.append("optimize_construct_or_display_strategy_first")
        else:
            next_actions.append("pause_main_screen_and_resolve_risk_or_accessibility")
        if normal_risk in {"unknown", "high"}:
            next_actions.append("complete_normal_tissue_review")
        if mixed_or_secreted:
            next_actions.append("confirm_cell_surface_exposure_by_flow_or_surface_biotinylation")

        ranked.append(
            {
                "rank_bucket": bucket,
                "rank_bucket_label": bucket_label,
                "report_priority_score": f"{score:.1f}",
                "identity_preference": identity_preference,
                "recommendation": recommendation,
                "candidate_id": candidate_id,
                "gene_symbol": gene_symbol,
                "uniprot_accession": row.get("uniprot_accession", ""),
                "protein_name": row.get("protein_name", ""),
                "membrane_group": join.get("membrane_group", ""),
                "antigen_accessibility": accessibility,
                "accessibility_gate": access_gate,
                "ecd_region": row.get("ecd_region", ""),
                "ecd_length": "" if ecd_length is None else str(ecd_length),
                "human_mouse_ecd_identity": "" if identity is None else f"{identity:.3f}",
                "human_mouse_ecd_identity_pct": _pct(identity),
                "human_mouse_divergence_class": divergence_class,
                "original_mouse_model_transferability_gate": row.get("mouse_model_transferability_gate", ""),
                "constructability_gate": construct_gate,
                "constructability_class": row.get("constructability_class", ""),
                "construct_recommendation": construct.get("construct_recommendation", ""),
                "primary_strategy": strategy.get("primary_strategy", ""),
                "secondary_strategy": strategy.get("secondary_strategy", ""),
                "normal_tissue_risk": normal_risk,
                "normal_tissue_risk_detail": high_risk_tissues,
                "pipeline_priority_call": row.get("priority_call", ""),
                "local_evidence_source_count": str(len(local_sources)),
                "local_evidence_sources": ";".join(local_sources),
                "risk_flags_original": ";".join(risk_flags),
                "readiness_flags": ";".join(readiness_flags),
                "passed_reasons": ";".join(passed_reasons),
                "caution_or_fail_reasons": ";".join(cautions) if cautions else "no_major_automated_caution",
                "next_action": ";".join(next_actions),
                "coverage_missing_reason": coverage.get("missing_reason_primary", ""),
            }
        )

    bucket_order = {"A": 0, "B": 1, "C": 2, "D": 3}
    ranked.sort(
        key=lambda item: (
            bucket_order.get(str(item["rank_bucket"]), 9),
            -float(str(item["report_priority_score"])),
            _safe_float(item["human_mouse_ecd_identity"]) if item["human_mouse_ecd_identity"] else 9.0,
            str(item["gene_symbol"]),
        )
    )
    for index, item in enumerate(ranked, start=1):
        item["report_rank"] = str(index)
    return ranked


def write_report_priority_outputs(
    ranked_rows: list[dict[str, Any]],
    outdir: Path,
    *,
    identity_preference: str = "low_for_antibody_screening",
    write_xlsx: bool = True,
) -> dict[str, Any]:
    outdir.mkdir(parents=True, exist_ok=True)
    table_path = outdir / "report_priority_table.tsv"
    with table_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=REPORT_PRIORITY_COLUMNS, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(ranked_rows)

    summary = {
        "status": "pass",
        "identity_preference": identity_preference,
        "candidate_count": len(ranked_rows),
        "rank_bucket_counts": dict(Counter(row["rank_bucket_label"] for row in ranked_rows)),
        "divergence_class_counts": dict(Counter(row["human_mouse_divergence_class"] for row in ranked_rows)),
        "normal_tissue_counts": dict(Counter(row["normal_tissue_risk"] for row in ranked_rows)),
        "outputs": {
            "report_priority_table": str(table_path),
            "report_priority_summary": str(outdir / "report_priority_summary.json"),
            "report_priority_report": str(outdir / "report_priority_report.md"),
        },
    }

    report_lines = [
        "# Report Priority Summary",
        "",
        f"Identity preference: `{identity_preference}`",
        f"Candidate count: {len(ranked_rows)}",
        "",
        "## Rank Buckets",
        *(f"- {key}: {value}" for key, value in sorted(summary["rank_bucket_counts"].items())),
        "",
        "## Human-Mouse ECD Divergence",
        *(f"- {key}: {value}" for key, value in sorted(summary["divergence_class_counts"].items())),
        "",
        "## Top 20",
        "| Rank | Gene | Bucket | ECD identity | Divergence | Accessibility | Construct | Main cautions |",
        "|---:|---|---|---:|---|---|---|---|",
    ]
    for row in ranked_rows[:20]:
        report_lines.append(
            "| {report_rank} | {gene_symbol} | {rank_bucket_label} | {human_mouse_ecd_identity_pct} | "
            "{human_mouse_divergence_class} | {antigen_accessibility} | {constructability_gate} | "
            "{caution_or_fail_reasons} |".format(**row)
        )
    report_path = outdir / "report_priority_report.md"
    report_path.write_text("\n".join(report_lines) + "\n", encoding="utf-8")

    if write_xlsx:
        xlsx_path = outdir / "report_priority_workbook.xlsx"
        _write_xlsx(ranked_rows, xlsx_path, summary)
        summary["outputs"]["report_priority_workbook"] = str(xlsx_path)
    (outdir / "report_priority_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return summary


def _write_xlsx(ranked_rows: list[dict[str, Any]], path: Path, summary: dict[str, Any]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError("openpyxl is required to write report_priority_workbook.xlsx") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["metric", "value"])
    ws.append(["identity_preference", summary["identity_preference"]])
    ws.append(["candidate_count", summary["candidate_count"]])
    ws.append(["rank_bucket_counts", json.dumps(summary["rank_bucket_counts"], sort_keys=True)])
    ws.append(["divergence_class_counts", json.dumps(summary["divergence_class_counts"], sort_keys=True)])
    ws.append(["important_note", "low_for_antibody_screening treats low human-mouse ECD identity as favorable; do not confuse it with mouse-model transferability."])

    for sheet_name, rows in [
        ("Presentation_ranked", ranked_rows),
        ("Top_50", ranked_rows[:50]),
        ("A_priority", [row for row in ranked_rows if row["rank_bucket"] == "A"]),
        ("B_backup_verify", [row for row in ranked_rows if row["rank_bucket"] == "B"]),
        ("C_conditional", [row for row in ranked_rows if row["rank_bucket"] == "C"]),
        ("D_deprioritize", [row for row in ranked_rows if row["rank_bucket"] == "D"]),
    ]:
        ws = wb.create_sheet(sheet_name)
        ws.append(REPORT_PRIORITY_COLUMNS)
        for row in rows:
            ws.append([row.get(column, "") for column in REPORT_PRIORITY_COLUMNS])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    fills = {
        "A": PatternFill("solid", fgColor="C6EFCE"),
        "B": PatternFill("solid", fgColor="D9EAD3"),
        "C": PatternFill("solid", fgColor="FFF2CC"),
        "D": PatternFill("solid", fgColor="F4CCCC"),
    }
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        headers = [cell.value for cell in ws[1]]
        bucket_index = headers.index("rank_bucket") + 1 if "rank_bucket" in headers else None
        for row_index in range(2, ws.max_row + 1):
            fill = fills.get(str(ws.cell(row_index, bucket_index).value or "")) if bucket_index else None
            for col_index in range(1, ws.max_column + 1):
                cell = ws.cell(row_index, col_index)
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if fill:
                    cell.fill = fill
        for col_index in range(1, ws.max_column + 1):
            width = 10
            for row_index in range(1, min(ws.max_row, 120) + 1):
                width = max(width, min(len(str(ws.cell(row_index, col_index).value or "")), 60))
            ws.column_dimensions[get_column_letter(col_index)].width = min(width + 2, 52)
    wb.save(path)


def run_report_priority(
    *,
    screening_table: Path,
    outdir: Path,
    readiness_table: Path | None = None,
    normal_tissue_risk: Path | None = None,
    construct_plan: Path | None = None,
    screening_strategy: Path | None = None,
    coverage_diagnostics: Path | None = None,
    join_evidence: Path | None = None,
    identity_preference: str = "low_for_antibody_screening",
    write_xlsx: bool = True,
) -> dict[str, Any]:
    screening_rows = _read_optional(screening_table)
    if not screening_rows:
        raise ValueError(f"screening table has no rows: {screening_table}")
    ranked_rows = rank_candidates(
        screening_rows,
        readiness_rows=_read_optional(readiness_table),
        normal_tissue_rows=_read_optional(normal_tissue_risk),
        construct_rows=_read_optional(construct_plan),
        strategy_rows=_read_optional(screening_strategy),
        coverage_rows=_read_optional(coverage_diagnostics),
        join_rows=_read_optional(join_evidence),
        identity_preference=identity_preference,
    )
    return write_report_priority_outputs(
        ranked_rows,
        outdir,
        identity_preference=identity_preference,
        write_xlsx=write_xlsx,
    )
